# -*- coding: utf-8 -*-
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import gwps_engine as E

tmp = tempfile.mkdtemp(prefix="gwps_footer_")
inp = os.path.join(tmp, "in.docx")
outp = os.path.join(tmp, "out.docx")

d = Document()
p = d.add_paragraph("关于开展××工作的通知"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph("各有关单位：")
d.add_paragraph("为深入推进相关工作，现将有关事项通知如下。")
d.add_paragraph("一、总体要求")
d.add_paragraph("1.细化措施")
d.add_paragraph("某某市人民政府办公室")
d.add_paragraph("2026年08月05日")
d.save(inp)

cfg = {
    "cc": "各县（区）人民政府",
    "print_org": "某某市人民政府办公室",
    "print_date": "2026年10月3日",
    "h3_bold": True,
}
proc = E.WordProcessor(config=cfg, log=lambda m: None)
report = proc.format_document(inp, outp)
rmap = {n: ok for (n, ok, _) in report}

fails = []
def chk(cond, msg):
    if not cond: fails.append(msg)

chk(rmap.get("版记(抄送/印发)(B3)") is True, "footer_note_check_fail:" + repr(rmap))
d2 = Document(outp)
all_text = "\n".join(p.text for p in d2.paragraphs)
assert "抄送：各县（区）人民政府。" in all_text, all_text
assert "某某市人民政府办公室" in all_text and "2026年10月3日印发" in all_text, all_text

# B5 三级标题加粗
h3_bold_ok = False
for p in d2.paragraphs:
    if p.text.strip().startswith("1."):
        for r in p.runs:
            if r.font.bold is not True:
                fails.append("h3_not_bold:" + repr(r.font.bold))
            else:
                h3_bold_ok = True
chk(h3_bold_ok, "h3_no_runs")

# 版记段落存在 pBdr（粗/细线）
border_ok = False
for p in d2.paragraphs:
    if not (p.text or "").strip():
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn("w:pBdr")) is not None:
            border_ok = True
            break
chk(border_ok, "no_borders")

if fails:
    print("FOOTER_FAIL")
    for f in fails:
        print("  ", f)
else:
    print("FOOTER_OK")
