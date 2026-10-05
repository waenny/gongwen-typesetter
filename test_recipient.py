# -*- coding: utf-8 -*-
"""主送机关识别回归：应顶格（firstLineChars=0），且主标题识别正常。"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import gwps_engine as E


def build_doc(path):
    d = Document()
    p = d.add_paragraph("关于加强安全生产工作的通知")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("各县（区）人民政府、各有关部门：")
    d.add_paragraph("为切实做好安全生产工作，现将有关事项通知如下。")
    d.add_paragraph("一、总体要求")
    d.add_paragraph("（一）提高思想认识")
    d.save(path)


def recipient_indent(doc):
    for p in doc.paragraphs:
        if p.text.strip().endswith("：") and "人民政府" in p.text:
            pPr = p._p.pPr
            if pPr is None:
                return "NOPPR"
            ind = pPr.find(qn("w:ind"))
            if ind is None:
                return "NOIND"
            return ind.get(qn("w:firstLineChars"))
    return "NOTFOUND"


tmp = tempfile.mkdtemp(prefix="gwps_rcp_")
inp = os.path.join(tmp, "in.docx")
outp = os.path.join(tmp, "out.docx")
build_doc(inp)

fails = []
proc = E.WordProcessor(config={}, log=lambda m: None)
rep = proc.format_document(inp, outp)
rmap = {n: ok for (n, ok, _) in rep}
if rmap.get("主标题识别") is not True:
    fails.append("title_fail")
v = recipient_indent(Document(outp))
if v not in ("0", None):
    fails.append("recipient_not_flush:" + str(v))

if fails:
    print("RECIPIENT_FAIL")
    for f in fails:
        print("  ", f)
else:
    print("RECIPIENT_OK")
