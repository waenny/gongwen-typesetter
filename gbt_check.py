# -*- coding: utf-8 -*-
"""GB/T 9704 公文排版自动核查器：对助手输出的 .docx 逐项校验，返回违规列表。"""
import sys
import re
from docx import Document
from docx.oxml.ns import qn

BODY = "仿宋_GB2312"
TITLE = "方正小标宋简体"
H1 = "黑体"
H2 = "楷体_GB2312"

RE_H1 = re.compile(r"^[一二三四五六七八九十]+、")
RE_H2 = re.compile(r"^[（(][一二三四五六七八九十]+[）)]")
RE_H3 = re.compile(r"^\d+\s*[.．]")
RE_H4 = re.compile(r"^[（(]\d+[）)]")
RE_DATE = re.compile(r"^\d{4}年\d{1,2}月\d{1,2}日$")
CJK_PUNCT = set("“”‘’、。，；：！？（）《》【】…—")
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_IMAGE_XPATH = ".//w:drawing|.//w:pict|.//w:object"


def rfonts(run):
    rf = run._r.find(qn("w:rPr"))
    if rf is None:
        return {}
    e = rf.find(qn("w:rFonts"))
    if e is None:
        return {}
    return {"eastAsia": e.get(qn("w:eastAsia")), "ascii": e.get(qn("w:ascii")),
            "hAnsi": e.get(qn("w:hAnsi")), "cs": e.get(qn("w:cs"))}


def rsize(run):
    return (run.font.size.pt if run.font.size else None)


def rspacing(run):
    rf = run._r.find(qn("w:rPr"))
    if rf is None:
        return None
    e = rf.find(qn("w:spacing"))
    return e.get(qn("w:val")) if e is not None else None


def outline(p):
    pPr = p._p.pPr
    if pPr is None:
        return None
    o = pPr.find(qn("w:outlineLvl"))
    return o.get(qn("w:val")) if o is not None else None


def firstline(p):
    pPr = p._p.pPr
    if pPr is None:
        return None
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        return None
    return ind.get(qn("w:firstLineChars"))


def check(path):
    d = Document(path)
    v = []
    sec = d.sections[0]
    if round(sec.page_width.cm, 1) != 21.0 or round(sec.page_height.cm, 1) != 29.7:
        v.append("页面非A4")
    if abs(sec.top_margin.cm - 3.7) > 0.1:
        v.append("天头≠37mm(%s)" % round(sec.top_margin.cm, 2))
    if abs(sec.left_margin.cm - 2.8) > 0.1:
        v.append("订口≠28mm(%s)" % round(sec.left_margin.cm, 2))
    if sec._sectPr.find(qn("w:docGrid")) is None:
        v.append("缺docGrid")

    alltext = "\n".join(p.text for p in d.paragraphs)
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells:
                alltext += "\n" + c.text
    if '"' in alltext or "'" in alltext:
        v.append("残留直引号 \" '")
    for p in d.paragraphs:
        t = p.text.strip()
        if re.match(r"^\d{4}年\d{1,2}月\d{1,2}日$", t) and ("年0" in t or "月0" in t):
            v.append("成文日期含虚位: " + t)

    # ★问题1★ 字间距：任何 run 的字符间距 w:spacing 非0即违规（正文+表格）
    def _paras_all():
        for p in d.paragraphs:
            yield p
        for tb in d.tables:
            for row in tb.rows:
                for c in row.cells:
                    for p in c.paragraphs:
                        yield p
    for p in _paras_all():
        for r in p.runs:
            sv = rspacing(r)
            if sv not in (None, "", "0"):
                v.append("字间距未归零: " + (p.text or "")[:12])
                break
    for p in d.paragraphs:
        if "DISTRIBUTE" in str(p.alignment):
            v.append("段落为分散对齐(应两端对齐): " + (p.text or "")[:12])
            break

    # ★字距被撑大★ 两端对齐的段落里若有软回车(w:br)，被软回车结束的那一行不算"段落最后一行"，
    # 会被撑满整行，中文就表现为字距被拉大（实测 16pt→56pt）；附件说明的 "1." 也会被推右而错位。
    for p in _paras_all():
        pPr = p._p.pPr
        e = pPr.find(qn("w:jc")) if pPr is not None else None
        jcv = e.get(qn("w:val")) if e is not None else None
        soft = [b for b in p._p.findall(".//" + qn("w:br"))
                if (b.get(qn("w:type")) or "") not in ("page", "column")]
        if jcv in ("both", "distribute") and soft:
            v.append("两端对齐+段内软回车(字距会被撑开): " + (p.text or "").replace("\n", " ")[:14])
            break

    def first_ea(p):
        for r in p.runs:
            e = rfonts(r).get("eastAsia")
            if e:
                return e
        return None

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        ea = first_ea(p)
        # 引号所在 run 的 hAnsi 不能是西文字体
        for r in p.runs:
            if any(ch in CJK_PUNCT for ch in r.text):
                if rfonts(r).get("hAnsi") == "Times New Roman":
                    v.append("中文标点按西文字形: " + t[:14])
                    break
        if RE_H1.match(t):
            if ea != H1:
                v.append("一级非黑体: " + t[:12])
            if outline(p) != "0":
                v.append("一级无大纲: " + t[:12])
            if p.runs and p.runs[0].font.bold:
                v.append("一级被加粗: " + t[:12])
        elif RE_H2.match(t):
            if ea != H2:
                v.append("二级非楷体: " + t[:12])
            if outline(p) != "1":
                v.append("二级无大纲: " + t[:12])
            if p.runs and p.runs[0].font.bold:
                v.append("二级被加粗: " + t[:12])
        elif RE_H3.match(t):
            if ea != BODY:
                v.append("三级非仿宋: " + t[:12])
            if outline(p) != "2":
                v.append("三级无大纲: " + t[:12])
        elif RE_H4.match(t):
            if ea != BODY:
                v.append("四级非仿宋: " + t[:12])
            if outline(p) != "3":
                v.append("四级无大纲: " + t[:12])
        elif re.match(r"^附件\d*$", t):
            if ea != H1:
                v.append("附件标记非黑体: " + t[:10])
            if p.runs and p.runs[0].font.bold:
                v.append("附件标记被加粗: " + t[:10])
        elif re.match(r"^附件[:：]", t):
            if ea not in (BODY, None):
                v.append("附件说明非仿宋: " + t[:12])
            idxs = [mm.start() for mm in re.finditer(r"\d+\.", t)]
            if len(idxs) >= 2:
                if "<w:br" not in p._p.xml:
                    v.append("多附件每项应换行: " + t[:12])
                ind = p._p.pPr.find(qn("w:ind")) if p._p.pPr is not None else None
                lc = ind.get(qn("w:leftChars")) if ind is not None else None
                if not lc or int(lc) < 400:
                    v.append("附件序号未悬挂对齐: " + t[:12])
        elif RE_DATE.match(t):
            if "RIGHT" not in str(p.alignment):
                v.append("成文日期未右对齐")
        elif t.endswith("：") and len(t) <= 60 and not RE_H1.match(t):
            fl = firstline(p)
            if fl not in (None, "0"):
                v.append("主送未顶格: " + t[:12])
        else:
            if t.endswith(("。", "！", "？")) and len(t) >= 15:  # 正文段
                if ea not in (BODY, None):
                    v.append("正文非仿宋: " + t[:12])
                if rsize(p.runs[0]) not in (16.0,) and p.runs:
                    v.append("正文非三号: " + t[:12])
                fl = firstline(p)
                if fl not in ("200", None):
                    v.append("正文首行未缩进2字: " + t[:12])
                spg = p._p.pPr.find(qn("w:snapToGrid")) if p._p.pPr is not None else None
                if spg is None or spg.get(qn("w:val")) not in ("false", "0"):
                    v.append("正文仍对齐网格(字距会被撑大): " + t[:12])

    # 引号开合方向/配对（逐段扫描）
    for p in d.paragraphs:
        st = 0
        for ch in p.text:
            if ch == "“":
                if st == 1:
                    v.append("引号叠开(方向错): " + p.text[:14]); break
                st = 1
            elif ch == "”":
                if st == 0:
                    v.append("引号先闭(方向错): " + p.text[:14]); break
                st = 0
    if "..." in alltext:
        v.append("残留省略号...")
    if "--" in alltext:
        v.append("残留双连字符--")
    if re.search(r"\d：\d", alltext):
        v.append("时间被误转全角冒号")
    for p in d.paragraphs:
        if "  " in p.text:
            v.append("残留连续半角空格: " + p.text[:14]); break
    for p in _paras_all():
        bad = "".join(c for c in ("\t", "\u00a0", "\u200b", "\ufeff") if c in (p.text or ""))
        if bad:
            v.append("残留空白类字符(%s): %s" % (repr(bad), (p.text or "")[:14]))
            break
    firstp = next((p for p in d.paragraphs if p.text.strip()), None)
    if firstp is not None and re.match(r"^关于", firstp.text.strip()):
        if first_ea(firstp) != TITLE:
            v.append("大标题非小标宋")
        if "CENTER" not in str(firstp.alignment):
            v.append("大标题未居中")
        if firstp.runs and rsize(firstp.runs[0]) not in (22.0, 21.5, 22.5):
            v.append("大标题非二号: %s" % rsize(firstp.runs[0]))

    # 表格：表头黑体加粗，表体仿宋
    for tb in d.tables:
        for ri, row in enumerate(tb.rows):
            for c in row.cells:
                for p in c.paragraphs:
                    if not (p.text or "").strip() or not p.runs:
                        continue
                    ea = rfonts(p.runs[0]).get("eastAsia")
                    if ri == 0:
                        if ea != H1 or not p.runs[0].font.bold:
                            v.append("表头未黑体加粗: " + p.text[:8])
                    else:
                        if ea not in (BODY, None):
                            v.append("表体非仿宋: " + p.text[:8])

    # ★表格★ 整表居中 + 单元格上下左右居中（公文惯例；表头黑体加粗见上）
    def _first_bad_cell(tb):
        for row in tb.rows:
            for c in row.cells:
                tcPr = c._tc.tcPr
                va = tcPr.find(qn("w:vAlign")) if tcPr is not None else None
                if (va.get(qn("w:val")) if va is not None else None) != "center":
                    return c, "上下"
                for p in c.paragraphs:
                    if (p.text or "").strip() and "CENTER" not in str(p.alignment):
                        return c, "左右"
        return None, ""

    for ti, tb in enumerate(d.tables):
        tblPr = tb._tbl.tblPr
        e = tblPr.find(qn("w:jc")) if tblPr is not None else None
        if (e.get(qn("w:val")) if e is not None else None) != "center":
            v.append("表格未整体居中(表%d)" % ti)
        bad, which = _first_bad_cell(tb)
        if bad is not None:
            v.append("单元格未%s居中(表%d): %s" % (which, ti, (bad.text or "")[:10]))

    # ★图片★ 含图片的段落：居中、无首行缩进、不能是固定行距（会裁图）、宽度不超版心
    max_cx = int(round((21.0 - sec.left_margin.cm - sec.right_margin.cm) * 360000))
    for p in d.paragraphs:
        if not p._p.xpath(_IMAGE_XPATH):
            continue
        head = (p.text or "")[:10]
        if "CENTER" not in str(p.alignment):
            v.append("图片段未居中: " + head)
            break
        pPr = p._p.pPr
        ind = pPr.find(qn("w:ind")) if pPr is not None else None
        if ind is not None and ind.get(qn("w:firstLineChars")) not in (None, "0"):
            v.append("图片段仍有首行缩进: " + head)
            break
        sp = pPr.find(qn("w:spacing")) if pPr is not None else None
        if sp is not None and sp.get(qn("w:lineRule")) in ("exact", "atLeast"):
            v.append("图片段是固定行距(会把图裁掉): " + head)
            break
        over = False
        for e2 in p._p.iter(_WP + "extent"):
            try:
                if int(e2.get("cx")) > max_cx:
                    over = True
            except (TypeError, ValueError):
                pass
        if over:
            v.append("图片宽超出\u7248\u5fc3: " + head)
            break
    return v


if __name__ == "__main__":
    violations = check(sys.argv[1])
    if violations:
        print("VIOLATIONS(%d):" % len(violations))
        for x in violations[:40]:
            print("  -", x)
    else:
        print("GBT_CHECK_OK")
