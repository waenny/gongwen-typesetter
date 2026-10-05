# -*- coding: utf-8 -*-
import sys, os, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gwps_engine as E
from docx import Document
from docx.oxml.ns import qn

inp, outp = sys.argv[1], sys.argv[2]
rep = E.WordProcessor(config={}, log=lambda m: None).format_document(inp, outp)
d = Document(outp)


def rattr(r, attr):
    rf = r._r.find(qn("w:rPr"))
    if rf is None:
        return None
    e = rf.find(qn("w:rFonts"))
    return e.get(qn(attr)) if e is not None else None


fails = []

# 1) 成文日期段落右对齐
dates = [p for p in d.paragraphs if re.match(r"^\d{4}年\d{1,2}月\d{1,2}日$", p.text.strip())]
if not dates:
    fails.append("no成文日期段")
for p in dates:
    if "RIGHT" not in str(p.alignment):
        fails.append("成文日期未右对齐:" + p.text)

# 2) 署名右对齐（结尾是委员会/办公室且短）
sigs = [p for p in d.paragraphs if p.text.strip().endswith(("委员会", "办公室")) and len(p.text.strip()) <= 30]
for p in sigs:
    if "RIGHT" not in str(p.alignment):
        fails.append("署名未右对齐:" + p.text.strip()[:12])

# 3) 中文标点走中文字形：取含中文标点的正文run，看 hAnsi
probe = None
for p in d.paragraphs:
    if "生命至上" in p.text and p.runs:
        probe = p.runs[0]
        break
if probe is not None:
    h = rattr(probe, "w:hAnsi")
    if h != "仿宋_GB2312":
        fails.append("hAnsi非中文字形:" + str(h))

# 4) 附件另页标记=黑体顶格不加粗
for p in d.paragraphs:
    if re.match(r"^附件\d*$", p.text.strip()):
        if p.runs:
            ea = rattr(p.runs[0], "w:eastAsia")
            if ea != "黑体":
                fails.append("附件标记非黑体:" + p.text.strip())
            if p.runs[0].font.bold:
                fails.append("附件标记被加粗:" + p.text.strip())

print("TITLE=", E.WordProcessor.__dict__ and "")
print("CHECKS=", len(rep))
print("SELFCHECK_OK" if not fails else "SELFCHECK_FAIL " + " | ".join(fails))
