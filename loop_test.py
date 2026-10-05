# -*- coding: utf-8 -*-
import os
import sys
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import gwps_engine as E
import gbt_check
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

CASES = os.path.join(HERE, "_cases")
shutil.rmtree(CASES, ignore_errors=True)
os.makedirs(CASES, exist_ok=True)


def mk(name, blocks):
    d = Document()
    for item in blocks:
        if item[0] == "P":
            p = d.add_paragraph()
            r = p.add_run(item[1])
            if len(item) > 2 and item[2] == "c":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif item[0] == "BR":  # 软回车
            p = d.add_paragraph()
            r = p.add_run(item[1]); r.add_break(); p.add_run(item[2])
        elif item[0] == "TBL":
            t = d.add_table(rows=len(item[1]), cols=len(item[1][0]))
            for i, row in enumerate(item[1]):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
    in_path = os.path.join(CASES, name + ".docx")
    d.save(in_path)
    return in_path


BODY1 = "为扎实做好安全生产工作，切实维护人民群众生命财产安全，根据有关法律法规，结合本地实际，现就有关事项通知如下。"
BODY2 = "各地区各有关部门要切实提高政治站位，压实责任，细化措施，狠抓落实，坚决防范遏制各类事故发生，确保形势持续稳定向好发展。"

cases = {}
cases["标准通知"] = [
    ("P", "　　进一步做好安全生产工作的通知。"),
    ("P", "各有关单位:"),
    ("P", BODY1),
    ("P", "一、总体要求"),
    ("P", "（一）指导思想"),
    ("P", BODY2),
    ("P", "1.近期目标"),
    ("P", "（1）完成隐患排查。"),
    ("P", "某某市人民政府办公室"),
    ("P", "2026年08月05日"),
]
cases["引导句多"] = [
    ("P", "安全生产专项整治方案"),
    ("P", "各县区:"),
    ("P", "一、总体要求"),
    ("P", "（一）提高认识。各单位要充分认识安全生产的重要性，切实增强责任感和紧迫感，把各项工作落到实处。"),
    ("P", "（二）压实责任。要明确职责分工，强化协同配合，形成齐抓共管的良好工作局面，确保取得实效。"),
    ("P", "二、工作任务"),
    ("P", "1.隐患排查"),
    ("P", "（1）全面排查各类风险隐患，建立台账，限期整改到位。"),
    ("P", "某某市人民政府办公室"),
    ("P", "2026年8月5日"),
]
cases["含表格"] = [
    ("P", "关于开展专项检查的通知"),
    ("P", "各单位:"),
    ("P", BODY1),
    ("P", "一、检查安排"),
    ("P", "本次检查分阶段推进,具体安排如下,"),
    ("TBL", [("阶段", "时间", "任务"), ("自查", "2026年8月", "自查自改,建立台账"), ("督查", "2026年9月", "抽查督办,落实整改")]),
    ("P", "某某市办公室"),
    ("P", "2026年08月05日"),
]
cases["标点极端"] = [
    ("P", "关于测试标点符号的通知"),
    ("P", "各单位:"),
    ("P", "他说\"很重要\",又提到'关键';还有等待中...区间2020--2025,电话0571-8888-1234,时间12:30,比例3:2,网址http://10.0.0.1:82。"),
    ("P", "第一，常态化排查\n对重点区域开展不间断巡查，发现问题立行立改。"),
    ("P", "二、附则。本通知自印发之日起施行。"),
    ("P", "某某市办公室"),
    ("P", "2026年08月05日"),
]
cases["极简"] = [("P", "_only title"), ("P", "关于某某工作的批复"), ("P", "某某市办公室"), ("P", "2026年08月05日")]

summary = []
for name, blocks in cases.items():
    inp = mk(name, blocks)
    outp = os.path.join(CASES, name + "_out.docx")
    try:
        E.WordProcessor(config={}, log=lambda m: None).format_document(inp, outp)
        viol = gbt_check.check(outp)
    except Exception as e:
        viol = ["EXCEPTION: " + repr(e)]
    summary.append((name, viol))
    print("\n==== 用例:", name, "→", "OK" if not viol else ("%d 违规" % len(viol)))
    for x in viol[:20]:
        print("   -", x)

total = sum(len(v) for _, v in summary)
print("\n#### 总计违规:", total, "####")
