# -*- coding: utf-8 -*-
import sys
from docx import Document
d = Document(sys.argv[1])
ts = [x.text for x in d.paragraphs]
tbl = []
for tb in d.tables:
    for row in tb.rows:
        for cell in row.cells:
            tbl.append(cell.text)
def has(s, pool):
    return any(s in t for t in pool)
print("表格仍含半角逗号:", any("," in t for t in tbl))
print("表格'制定方案，发动部署':", has("制定方案，发动部署", tbl))
print("标题以'关于'开头:", any(t.strip().startswith("关于") for t in ts))
print("落款去虚位2026年8月5日:", sum(t.strip() == "2026年8月5日" for t in ts))
print("正文电话保留0571-8888-1234:", has("0571-8888-1234", ts))
print("引号“要注意”:", has("“要注意”", ts))
