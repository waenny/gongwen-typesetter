# -*- coding: utf-8 -*-
import os
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import gwps_engine as E

tmp_in = os.path.join(HERE, "_proof_in.docx")
tmp_out = os.path.join(HERE, "_proof_out.docx")

doc = Document()

def add_sp(rPr, val="60"):
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), val)
    rPr.append(sp)

# 标题
doc.add_paragraph("关于测试字符间距的通知")
# 正文：run 上设字间距
p = doc.add_paragraph("这是一段把字符间距加宽到60的正文，应当被清零。")
add_sp(p.runs[0]._r.get_or_add_rPr(), "60")
# 段落标记(pPr/rPr)上设字间距
p2 = doc.add_paragraph("这段把间距设在段落标记上，也应被清零。")
rPr = p2._p.get_or_add_pPr()
mark = OxmlElement("w:rPr")
add_sp(mark, "80")
rPr.append(mark)
# Normal 样式上设字间距
nrm = doc.styles["Normal"]._element.get_or_add_rPr()
add_sp(nrm, "40")
doc.add_paragraph("附件：1.安全生产检查表2.隐患整改清单")
doc.save(tmp_in)

E.WordProcessor(config={}, log=lambda m: None).format_document(tmp_in, tmp_out)

from docx.oxml import parse_xml
import zipfile
with zipfile.ZipFile(tmp_out) as z:
    docxml = z.read("word/document.xml").decode("utf-8", "ignore")
    try:
        stx = z.read("word/styles.xml").decode("utf-8", "ignore")
    except KeyError:
        stx = ""
import re as _re
_cs = lambda s: len(_re.findall(r"<w:spacing[^>]*w:val=", s))
_ww = lambda s: len(_re.findall(r"<w:w [^>]*", s))
print("document.xml 残留字符间距(w:spacing@val):", _cs(docxml))
print("styles.xml   残留字符间距(w:spacing@val):", _cs(stx))
print("残留字符缩放 w:w 个数:", _ww(docxml) + _ww(stx))

d2 = Document(tmp_out)
for p in d2.paragraphs:
    if p.text.strip().startswith("附件："):
        print("附件说明段 XML:")
        print("   含<w:br>:", "<w:br" in p._p.xml)
        ind = p._p.pPr.find(qn("w:ind")) if p._p.pPr is not None else None
        print("   leftChars=", ind.get(qn("w:leftChars")) if ind is not None else None,
              "hangingChars=", ind.get(qn("w:hangingChars")) if ind is not None else None)
        print("   各行:", p.text.split("\n") if p.text else [])
        break
