# -*- coding: utf-8 -*-
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor
from docx.oxml.ns import qn
import gwps_engine as E

tmp = tempfile.mkdtemp(prefix="gwps_hdr_")
inp = os.path.join(tmp, "in.docx"); outp = os.path.join(tmp, "out.docx")

d = Document()
p = d.add_paragraph("关于开展安全生产大检查的通知"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph("各有关单位：")
d.add_paragraph("现将有关事项通知如下。")
d.add_paragraph("一、总体要求")
d.add_paragraph("某某市人民政府")
d.add_paragraph("2026年08月05日")
d.save(inp)

cfg = {
    "header_style": "正式发文（大红头）",
    "org_name": "某某市人民政府",
    "doc_number": "某政发〔2026〕5号",
    "signer": "张三",
}
proc = E.WordProcessor(config=cfg, log=lambda m: None)
report = proc.format_document(inp, outp)
fails = []

# 主标题仍是"关于…"，不是机关名版头
if not proc.last_title.startswith("关于"):
    fails.append("title_is_mast:" + proc.last_title)

d2 = Document(outp)
paras = [p.text.strip() for p in d2.paragraphs]
joined = "\n".join(paras)

# 大红头机关标志存在
if "某某市人民政府文件" not in joined:
    fails.append("mast_missing")
# 签发人整行存在，姓名楷体
signer_run_kaiti = False
org_run_red = False
for p in d2.paragraphs:
    if "签发人：" in p.text and "张三" in p.text:
        for r in p.runs:
            rf = r._r.find(qn("w:rPr"))
            ea = rf.find(qn("w:rFonts")).get(qn("w:eastAsia")) if (rf is not None and rf.find(qn("w:rFonts")) is not None) else None
            if r.text == "张三" and ea == "楷体_GB2312":
                signer_run_kaiti = True
    if p.text.strip() == "某某市人民政府文件":
        for r in p.runs:
            try:
                if r.font.color and r.font.color.rgb == RGBColor(210, 31, 43):
                    org_run_red = True
            except Exception:
                pass
if not signer_run_kaiti:
    fails.append("signer_not_kaiti_or_missing")
if not org_run_red:
    fails.append("mast_not_red")

# docGrid linePitch 解耦为 580
grid = d2.sections[0]._sectPr.find(qn("w:docGrid"))
if grid is None or grid.get(qn("w:linePitch")) != "580":
    fails.append("grid_pitch:" + (grid.get(qn("w:linePitch")) if grid is not None else "none"))

# 顶部留白：机关标志之前存在空段
idx_org = next((i for i, t in enumerate(paras) if t == "某某市人民政府文件"), -1)
blanks_before = sum(1 for t in paras[:idx_org] if t == "") if idx_org > 0 else 0
if blanks_before < 2:
    fails.append("no_top_blank_lines:%d" % blanks_before)

print("HEADER_OK" if not fails else "HEADER_FAIL " + "; ".join(fails))
