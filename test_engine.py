# -*- coding: utf-8 -*-
import os, sys, re, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import gwps_engine as E


def ea_runs(para):
    out = []
    for r in para.runs:
        e = None
        rf = r._r.find(qn("w:rPr"))
        if rf is not None:
            rfonts = rf.find(qn("w:rFonts"))
            if rfonts is not None:
                e = rfonts.get(qn("w:eastAsia"))
        out.append((r.text, e))
    return out


def find_para(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


tmp = tempfile.mkdtemp(prefix="gwps_smoke_")
inp = os.path.join(tmp, "in.docx")
outp = os.path.join(tmp, "out.docx")

d = Document()
p = d.add_paragraph("关于开展××工作的通知"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
d.add_paragraph("各有关单位：")
d.add_paragraph("为深入推进相关工作，现将有关事项通知如下。")
d.add_paragraph("一、总体要求")
d.add_paragraph("（一）加强组织领导")
d.add_paragraph("（二）压实责任。各单位要抓紧落实，确保成效明显，务求实效到位。")
d.add_paragraph("1.细化工作措施")
d.add_paragraph('通知指出"要高度重视"，请抓好落实。')
d.add_paragraph("某某市人民政府办公室")
d.add_paragraph("2026年08月05日")
d.save(inp)

proc = E.WordProcessor(log=lambda m: None)
report = proc.format_document(inp, outp)
rmap = {name: ok for (name, ok, detail) in report}

fails = []

def chk(cond, msg):
    if not cond:
        fails.append(msg)

chk(rmap.get("主标题识别") is True, "title_not_detected")
chk(rmap.get("22行/28字网格(docGrid)") is True, "docgrid_missing")
chk(rmap.get("正文字体 仿宋_GB2312(eastAsia)") is True, "body_eastAsia_missing")
chk(rmap.get("标题大纲级别(outlineLvl)") is True, "outline_missing")

d2 = Document(outp)
# A2: 独立（一）整行楷体
h2 = find_para(d2, "（一）")
if h2:
    runs = ea_runs(h2)
    chk(all(e == "楷体_GB2312" for _, e in runs if e) and any(e == "楷体_GB2312" for _, e in runs),
        "standalone_h2_not_all_kaiti:" + repr(runs))
else:
    fails.append("standalone_h2_missing")

# A2: 引导句（二）...。正文 → 同段既有楷体又有仿宋
lead = find_para(d2, "（二）")
if lead:
    fonts = set(e for _, e in ea_runs(lead) if e)
    chk("楷体_GB2312" in fonts and "仿宋_GB2312" in fonts,
        "leadin_split_bad:" + repr(sorted(fonts)))
else:
    fails.append("leadin_missing")

# 引号规范化为中文上下引号
alltext = "\n".join(p.text for p in d2.paragraphs)
chk("\u201c" in alltext and "\u201d" in alltext, "quotes_not_curly")

# 一级标题黑体
h1 = find_para(d2, "一、")
if h1:
    chk(all(e == "黑体" for _, e in ea_runs(h1) if e), "h1_not_heiti:" + repr(ea_runs(h1)))

# 成文日期去虚位
chk("2026年8月5日" in alltext, "date_not_normalized")

if fails:
    print("ENGINE_SMOKE_FAIL")
    for f in fails:
        print("  FAIL:", f)
else:
    print("ENGINE_SMOKE_OK  checks=" + str(len(report)))
print("REPORT:")
for name, ok, detail in report:
    print("  [%s] %s %s" % ("PASS" if ok else "FAIL", name, ("| "+detail) if detail else ""))
