# -*- coding: utf-8 -*-
"""排版引擎（重建版）—— WordProcessor

串起：导入转换 → 文本规范化(安全) → ★自动编号展开(A1)★ → 去空段(保媒体) → 版头三档 →
主/副标题识别(A3) → 逐段套用（标题层级字体+大纲、正文、★修正的二级标题切分(A2)★、
主送顶格、附件）→ 落款/成文日期 → 版记(B3) → 表格 → 页面/docGrid(B2)/单双页页码 → 自检(C2)。

python-docx 版式按 GB/T 9704；字体同时写 eastAsia/ascii/hAnsi/cs；OOXML 子元素按序插入。
"""
import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

import gwps_normalize as N
import gwps_structure as S
from gwps_config import (
    DEFAULT_CONFIG, HEADER_STYLE_NONE, HEADER_STYLE_BIG,
    BODY_FONT, TITLE_FONT, H1_FONT, H2_FONT, WESTERN_FONT, PAGE_NUMBER_FONT,
    MAST_RED_RGB,
)

# OOXML 顺序：把某元素插到"必须排在它之后"的元素前面
_AFTER_PBDR = ('w:shd', 'w:tabs', 'w:spacing', 'w:ind', 'w:jc', 'w:textAlignment',
               'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
_AFTER_OUTLINE = ('w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
_AFTER_DOCGRID = ('w:printerSettings', 'w:sectPrChange')

_MEDIA_XPATH = './/w:drawing|.//w:pict|.//w:object|.//w:txbxContent|.//w:hyperlink'
_BLANK_MEDIA_XPATH = ('.//w:drawing|.//w:pict|.//w:object|.//w:txbxContent'
                      '|.//w:fldSimple|.//w:hyperlink|.//w:br')
_ORG_TAIL = r"(委员会|办公室|人民政府|党委|党组|局|厅|部|处|科|院|中心|集团|公司|协会|学会|学校)$"

# 图片/图形段（注意不含 w:hyperlink，普通超链接段不该按图片处理）
_IMAGE_XPATH = './/w:drawing|.//w:pict|.//w:object'
# 图题：紧跟图片、以“图1 / 图一 / 图 1.”开头
_RE_FIGCAP = re.compile(r"^图\s*[0-9０-９一二三四五六七八九十]+\s*[\.、:：]?\s*\S")
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_WP_NS = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_EMU_PER_CM = 360000


def _cn_date(text):
    if not text:
        return text
    t = re.sub(r"年0(\d)", r"年\1", text)
    return re.sub(r"月0(\d)", r"月\1", t)


class WordProcessor:
    def __init__(self, config=None, log=None):
        self.config = dict(DEFAULT_CONFIG)
        if config:
            self.config.update(config)
        self.log = log or (lambda m: None)
        self.last_title = ""
        self._mast_ids = set()

    # ---------------- 导入：任意来源 → .docx ----------------
    def convert_to_docx(self, path):
        ext = os.path.splitext(path)[1].lower()
        if ext == ".docx":
            fd, tmp = tempfile.mkstemp(suffix=".docx", prefix="gwps_")
            os.close(fd)
            shutil.copy2(path, tmp)
            return tmp
        if ext in (".txt", ".md"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            if ext == ".md":
                content = self._clean_markdown(content)
            doc = Document()
            for line in content.splitlines() or [""]:
                doc.add_paragraph(line)
            fd, tmp = tempfile.mkstemp(suffix=".docx", prefix="gwps_")
            os.close(fd)
            doc.save(tmp)
            return tmp
        if ext in (".doc", ".wps"):
            return self._convert_legacy(path)
        raise ValueError("不支持的文件格式：%s" % ext)

    @staticmethod
    def _clean_markdown(text):
        text = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        text = re.sub(r"<[^>]+>", "", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
        text = re.sub(r"^\s*#+\s+", "", text, flags=re.M)
        text = re.sub(r"^\s*>\s?", "", text, flags=re.M)
        return re.sub(r"^\s*[-*_]{3,}\s*$", "", text, flags=re.M)

    def _convert_legacy(self, path):
        app = d = None
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            for prog in ("Word.Application", "KWPS.Application"):
                try:
                    app = win32com.client.DispatchEx(prog)
                    break
                except Exception:
                    app = None
            if app is not None:
                try:
                    app.Visible = False
                    app.DisplayAlerts = 0
                except Exception:
                    pass
                fd, tmp = tempfile.mkstemp(suffix=".docx", prefix="gwps_leg_")
                os.close(fd)
                d = app.Documents.Open(os.path.abspath(path))
                try:
                    d.SaveAs2(tmp, FileFormat=16)
                except Exception:
                    d.SaveAs(tmp, FileFormat=16)
                return tmp
        except Exception as e:
            self.log("  > WPS/Word 转换失败：%s" % e)
        finally:
            try:
                if d is not None:
                    d.Close(SaveChanges=0)
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            try:
                import pythoncom
                pythoncom.CoUninitialize()
            except Exception:
                pass
        soffice = self._find_soffice()
        if soffice:
            outdir = tempfile.mkdtemp(prefix="gwps_soffice_")
            kw = {"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {}
            try:
                subprocess.run([soffice, "--headless", "--norestore", "--convert-to",
                                "docx", "--outdir", outdir, os.path.abspath(path)],
                               timeout=180, **kw)
            except subprocess.TimeoutExpired:
                raise RuntimeError("LibreOffice 转换超时（>180s）。")
            res = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".docx")
            if os.path.isfile(res):
                return res
        raise RuntimeError("旧格式需已安装 Word/WPS 或 LibreOffice 才能转换。")

    @staticmethod
    def _find_soffice():
        for p in ("soffice", "soffice.exe"):
            r = shutil.which(p)
            if r:
                return r
        for env in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
            base = os.environ.get(env)
            if base:
                cand = os.path.join(base, "LibreOffice", "program", "soffice.exe")
                if os.path.isfile(cand):
                    return cand
        return None

    # ---------------- 底层工具 ----------------
    @staticmethod
    def _place(parent, child, successors):
        for e in parent.findall(child.tag):
            parent.remove(e)
        parent.insert_element_before(child, *successors)

    def _set_run_font(self, run, chinese_font, size_pt, color=None, bold=None, western_font=None):
        run.font.size = Pt(float(size_pt))
        if color is not None:
            run.font.color.rgb = color
        elif run.font.color and run.font.color.rgb is not None:
            try:
                rPr = run._r.get_or_add_rPr()
                old = rPr.find(qn("w:color"))
                if old is not None:
                    rPr.remove(old)
            except Exception:
                pass
        if bold is not None:
            run.font.bold = bold
        rFonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        for a in ("w:eastAsiaTheme", "w:asciiTheme", "w:hAnsiTheme", "w:csTheme", "w:cstheme"):
            rFonts.attrib.pop(qn(a), None)
        if western_font:
            en = western_font
        elif self.config.get("normalize_western", True):
            en = self.config.get("western_font", WESTERN_FONT)
        else:
            en = chinese_font
        rFonts.set(qn("w:eastAsia"), chinese_font)
        rFonts.set(qn("w:ascii"), en)
        # hAnsi/cs 用中文字体：使“ ” ‘ ’ 、。；：等在中文语境下走中文字形（避免按西文渲染导致间距/方向不对）
        rFonts.set(qn("w:hAnsi"), chinese_font)
        rFonts.set(qn("w:cs"), chinese_font)
        # ★问题1★ 归零字符间距：源文字若设了放大的字间距(w:rPr/w:spacing)，一律清除（正文/标题/表格/附件均适用）
        rPr = run._r.get_or_add_rPr()
        _sp = rPr.find(qn("w:spacing"))
        if _sp is not None:
            rPr.remove(_sp)

    def _apply_font_to_para(self, para, font, size, color=None, bold=None, western_font=None):
        for run in para.runs:
            self._set_run_font(run, font, size, color, bold, western_font)

    def _set_para_text_safe(self, para, new_text):
        """安全改文本：保留首 run 格式；含图片/超链接/公式的段跳过（返回 False）。"""
        if para._p.xpath(_MEDIA_XPATH):
            return False
        runs = para.runs
        if not runs:
            para.add_run(new_text)
            return True
        runs[0].text = new_text
        for r in runs[1:]:
            r.text = ""
        return True

    def _set_line_spacing(self, para, pt):
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(float(pt))

    def _set_first_line_indent(self, para, chars):
        ind = para._p.get_or_add_pPr().get_or_add_ind()
        ind.set(qn("w:firstLineChars"), str(int(chars) * 100))
        ind.set(qn("w:firstLine"), "0")

    def _set_left_indent(self, para, chars):
        ind = para._p.get_or_add_pPr().get_or_add_ind()
        ind.set(qn("w:leftChars"), str(int(chars) * 100))
        ind.set(qn("w:left"), "0")

    def _set_right_indent(self, para, chars):
        ind = para._p.get_or_add_pPr().get_or_add_ind()
        ind.set(qn("w:rightChars"), str(int(chars) * 100))
        ind.set(qn("w:right"), "0")

    def _clear_indent(self, para, also_left=False):
        ind = para._p.get_or_add_pPr().find(qn("w:ind"))
        if ind is None:
            return
        attrs = ["w:firstLineChars", "w:firstLine", "w:hangingChars", "w:hanging"]
        if also_left:
            attrs += ["w:left", "w:leftChars"]
        for a in attrs:
            ind.attrib.pop(qn(a), None)

    def _set_outline(self, para, level):
        if not self.config.get("set_outline", True):
            return
        o = OxmlElement("w:outlineLvl")
        o.set(qn("w:val"), str(int(level) - 1))
        self._place(para._p.get_or_add_pPr(), o, _AFTER_OUTLINE)

    def _para_border(self, para, top_sz=None, bottom_sz=None, color="000000"):
        pbdr = OxmlElement("w:pBdr")
        for side, sz in (("top", top_sz), ("bottom", bottom_sz)):
            if sz is None:
                continue
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "1")
            e.set(qn("w:color"), color)
            pbdr.append(e)
        self._place(para._p.get_or_add_pPr(), pbdr, _AFTER_PBDR)

    # ---------------- 逐段套用 ----------------
    def _format_body(self, para):
        self._apply_font_to_para(para, self.config.get("body_font", BODY_FONT),
                                 self.config.get("body_size", 16))
        self._set_first_line_indent(para, int(self.config.get("indent_chars", 2)))
        self._set_line_spacing(para, self.config.get("line_spacing", 28))
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def _format_recipient(self, para):
        """主送机关：仿宋三号、左顶格（清首行/左缩进/悬挂）。"""
        self._apply_font_to_para(para, self.config.get("body_font", BODY_FONT),
                                 self.config.get("body_size", 16))
        self._clear_indent(para, also_left=True)
        self._set_first_line_indent(para, 0)
        self._set_line_spacing(para, self.config.get("line_spacing", 28))
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _format_title(self, para):
        import re as _re
        cleaned = _re.sub(r"[。！？；，、：]+$", "", (para.text or "").strip())
        if cleaned and cleaned != para.text:
            self._set_para_text_safe(para, cleaned)
        self._apply_font_to_para(para, self.config.get("title_font", TITLE_FONT),
                                 self.config.get("title_size", 22))
        self._set_line_spacing(para, self.config.get("title_line_spacing", 33))
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._clear_indent(para, also_left=True)

    def _format_attachment_note(self, para):
        """附件说明：仿宋正文；多附件每项一行(软回车)+悬挂缩进使序号对齐；单附件左空二字。"""
        import re as _re
        t = (para.text or "").strip()
        m = _re.match(r"^(附件[:：])\s*(.*)$", t, _re.S)
        head = m.group(1) if m else "附件："
        items = []
        if m:
            rest = m.group(2).replace("\t", " ").replace("\n", " ")
            items = [x.strip() for x in _re.split(r"(?=\d+\.)", rest) if x.strip()]
        self._clear_indent(para, also_left=True)
        if len(items) >= 2:
            for r in list(para.runs):
                r._r.getparent().remove(r._r)
            parts = [head + items[0]] + items[1:]
            for i, seg in enumerate(parts):
                run = para.add_run(seg)
                if i < len(parts) - 1:
                    run.add_break()
            ind = para._p.get_or_add_pPr().get_or_add_ind()
            # 悬挂：行首“附件：”左空二字(2)，行首行缩回3字→起始在2；续行“N.”落在第5字→与“1.”对齐
            ind.set(qn("w:leftChars"), "500")
            ind.set(qn("w:left"), "0")
            ind.set(qn("w:hangingChars"), "300")
            ind.set(qn("w:hanging"), "0")
        else:
            self._set_left_indent(para, 2)
        self._apply_font_to_para(para, self.config.get("body_font", BODY_FONT),
                                 self.config.get("body_size", 16))
        self._set_line_spacing(para, self.config.get("line_spacing", 28))
        # ★序号对齐★ 这里必须左对齐：本段靠软回车实现"每项一行"，
        # 若设成两端对齐，第 1 行会被撑满整行，把 "1." 推右约 2 字，与第二行的 "2." 错开。
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def _strip_char_spacing(self, doc):
        """★问题1彻底★ 清除所有字符间距(w:spacing)与字符缩放(w:w)：正文/表格/段落标记/样式/页脚。"""
        roots = [doc.element]
        try:
            sel = getattr(doc.styles, "element", None)
            if sel is None:
                sel = getattr(doc.styles, "_element", None)
            if sel is not None:
                roots.append(sel)
        except Exception:
            pass
        for s in doc.sections:
            for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                       s.even_page_header, s.even_page_footer):
                try:
                    roots.append(hf._element)
                except Exception:
                    pass
        for root in roots:
            if root is None:
                continue
            for rPr in root.iter(qn("w:rPr")):
                for tag in ("w:spacing", "w:w"):
                    e = rPr.find(qn(tag))
                    if e is not None:
                        rPr.remove(e)

    def _disable_grid_snap(self, doc):
        """★字间距过大★ 关闭所有段落“对齐到文档网格”，防止字符网格把字距撑大（保留行距）。"""
        succ = ('w:spacing', 'w:ind', 'w:contextualSpacing',
                'w:jc', 'w:outlineLvl', 'w:rPr', 'w:sectPr', 'w:pPrChange')
        n = 0
        for pel in doc.element.iter(qn('w:p')):
            pPr = pel.find(qn('w:pPr'))
            if pPr is None:
                pPr = OxmlElement('w:pPr')
                pel.insert(0, pPr)
            e = pPr.find(qn('w:snapToGrid'))
            if e is None:
                e = OxmlElement('w:snapToGrid')
                self._place(pPr, e, succ)
            e.set(qn('w:val'), 'false')
            n += 1
        return n

    def _strip_soft_breaks(self, doc):
        """★字距过大★ 删除正文段落内的软回车(w:br/Shift+Enter)。

        Word 的两端对齐只放过"段落最后一行"；被软回车结束的那一行不算最后一行，
        于是被撑满整行，中文就表现为**字距被拉大**（实测 16pt 字距被拉到 56pt）。
        同时它把附件说明里的 "1." 推右，导致序号与第二行的 "2." 对不齐。
        有文字的段落一律合并回一段（等同于手工删掉那个 ↓），
        只含软回车的空段（附件另页留白）不动。
        注意：Word 把 .doc 转 .docx 时手动换行符写作 w:type="textWrapping"，
        python-docx 写作不带属性的 <w:br/>，两者都必须删；只放过分页/分栏符。
        """
        n = 0
        for p in doc.paragraphs:
            if not (p.text or "").strip():
                continue
            for br in p._p.findall(".//" + qn("w:br")):
                if (br.get(qn("w:type")) or "") in ("page", "column"):
                    continue          # 保留分页符/分栏符
                br.getparent().remove(br)
                n += 1
        return n

    def _format_subtitle(self, para):
        self._apply_font_to_para(para, self.config.get("subtitle_font", TITLE_FONT),
                                 self.config.get("subtitle_size", 16))
        self._set_line_spacing(para, self.config.get("title_line_spacing", 33))
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._clear_indent(para, also_left=True)

    def _format_heading(self, para, level):
        cfg = self.config
        if level == 1:
            font, size, bold = cfg.get("h1_font", H1_FONT), cfg.get("h1_size", 16), cfg.get("h1_bold")
        elif level == 2:
            font, size, bold = cfg.get("h2_font", H2_FONT), cfg.get("h2_size", 16), cfg.get("h2_bold")
        else:
            font, size = cfg.get("body_font", BODY_FONT), cfg.get("body_size", 16)
            bold = cfg.get("h3_bold") if level == 3 else cfg.get("h4_bold")
        self._apply_font_to_para(para, font, size, bold=bold)
        self._set_first_line_indent(para, int(cfg.get("indent_chars", 2)))
        self._set_line_spacing(para, cfg.get("line_spacing", 28))
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._set_outline(para, level)

    # ★A2 修正★：仅当确为"标题+同行正文(句号后有实质正文)"才切分
    def _try_h2_inline_body(self, para):
        if para._p.xpath(_MEDIA_XPATH):
            return False
        text = para.text or ""
        m = S._RE_H2.match(text.strip())
        if not m:
            return False
        head_end = len(m.group(0))
        rest = text[head_end:].lstrip()
        if not rest:
            return False
        mp = re.search(r"[。！？；]", rest)
        title_part = rest[:mp.start() + 1] if mp else rest
        body_part = rest[mp.end():] if mp else ""
        if mp is None or len(body_part.strip()) < 4:
            return False  # 独立二级标题 → 交 _format_heading(2) 整行楷体
        para._p.clear_content()
        r1 = para.add_run(text[:head_end] + title_part)
        self._set_run_font(r1, self.config.get("h2_font", H2_FONT), self.config.get("body_size", 16))
        if body_part:
            r2 = para.add_run(body_part)
            self._set_run_font(r2, self.config.get("body_font", BODY_FONT), self.config.get("body_size", 16))
        self._set_first_line_indent(para, int(self.config.get("indent_chars", 2)))
        self._set_line_spacing(para, self.config.get("line_spacing", 28))
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._set_outline(para, 2)
        return True

    def _is_attachment(self, para):
        return bool(re.match(r"^附件\s*[:：]", (para.text or "").strip()))

    # ---------------- 落款 / 日期 ----------------
    def _find_trailing_signature(self, blocks):
        """全文扫描：取最后一个“整段=成文日期”的段落及其上方紧邻的发文机关署名（不限文末）。"""
        paras = [b for b in blocks if isinstance(b, Paragraph)]
        date_para = None
        issuer = None
        for idx, p in enumerate(paras):
            if re.match(r"^\d{4}年\d{1,2}月(\d{1,2}日)?$", (p.text or "").strip()):
                date_para = p
                issuer = None
                for k in range(idx - 1, max(-1, idx - 3), -1):
                    pt = (paras[k].text or "").strip()
                    if not pt:
                        continue
                    if len(pt) <= 30 and not re.search(r"[。！？；，,]", pt) and re.search(_ORG_TAIL, pt):
                        issuer = paras[k]
                    break
        return date_para, issuer

    def _apply_sig_para(self, para, text=None):
        if text is not None:
            self._set_para_text_safe(para, text)
        self._apply_font_to_para(para, self.config.get("body_font", BODY_FONT),
                                 self.config.get("body_size", 16))
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._clear_indent(para, also_left=True)
        self._set_right_indent(para, int(self.config.get("sign_right_chars", 4)))

    def _format_signature(self, date_para, issuer_para):
        if issuer_para is not None:
            self._apply_sig_para(issuer_para)
        if date_para is not None:
            self._apply_sig_para(date_para, _cn_date(date_para.text.strip()))

    def _append_signature_block(self, doc):
        issuer = self.config.get("issuer", "")
        date_text = self.config.get("date_text", "")
        if not (issuer or date_text):
            return
        for _ in range(int(self.config.get("sign_blank", 1))):
            doc.add_paragraph("")
        if issuer:
            self._apply_sig_para(doc.add_paragraph(issuer))
        if date_text:
            self._apply_sig_para(doc.add_paragraph(_cn_date(date_text)))

    # ---------------- 版记（B3）----------------
    def _build_footer_notes(self, doc):
        c = self.config
        if not c.get("enable_footer_notes", True):
            return
        cc = (c.get("cc") or "").strip()
        org = (c.get("print_org") or "").strip()
        date = (c.get("print_date") or "").strip()
        if not (cc or org or date):
            return
        body = c.get("body_font", BODY_FONT)
        size = 14
        spacing = c.get("line_spacing", 28)

        def line(top=None, bottom=None):
            p = doc.add_paragraph("")
            self._para_border(p, top_sz=top, bottom_sz=bottom)
            self._set_line_spacing(p, 14)
            return p

        def elem(text, align, left=True, right=False):
            p = doc.add_paragraph()
            r = p.add_run(text)
            self._set_run_font(r, body, size)
            p.alignment = align
            if left:
                self._set_left_indent(p, 1)
            if right:
                self._set_right_indent(p, 1)
            self._set_line_spacing(p, spacing)
            return p

        line(top=8)
        if cc:
            elem("抄送：" + cc.rstrip("。；") + "。", WD_ALIGN_PARAGRAPH.LEFT)
        line(top=6)
        if org:
            elem(org, WD_ALIGN_PARAGRAPH.LEFT, right=False)
        if date:
            elem(date + "印发", WD_ALIGN_PARAGRAPH.RIGHT, left=False, right=True)
        line(bottom=8)
        self.log("  > 已插入版记(B3)。")

    # ---------------- 表格 ----------------
    def _format_tables(self, doc):
        """表格按公文惯例处理：整表居中；单元格**上下左右居中**；表头黑体加粗、表体仿宋。

        说明：GB/T 9704—2012 没有逐条规定单元格垂直居中，这是各级机关模板的通行做法
        （表头黑体居中、表内文字水平+垂直居中、表格整体居中）。
        """
        if not self.config.get("enable_table_formatting", True):
            return
        font = self.config.get("table_font", BODY_FONT)
        hfont = self.config.get("table_header_font", H1_FONT)
        size = self.config.get("table_size", 14)
        center_body = bool(self.config.get("table_body_center", True))
        for tb in doc.tables:
            try:
                tb.alignment = WD_TABLE_ALIGNMENT.CENTER      # 表格整体居中
            except Exception:
                pass
            for ri, row in enumerate(tb.rows):
                try:
                    row.height = Cm(self.config.get("table_row_height_cm", 0.8))
                    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                except Exception:
                    pass
                for cell in row.cells:
                    try:
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER   # 上下居中
                    except Exception:
                        pass
                    for para in cell.paragraphs:
                        head = (ri == 0)
                        self._apply_font_to_para(
                            para, hfont if head else font, size,
                            bold=True if (head and self.config.get("table_header_bold", True)) else None)
                        self._clear_indent(para, also_left=True)
                        # 左右居中：表头必居中；表体默认也居中（可在“选项”里关掉改成左对齐）
                        para.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                                          if (head or center_body) else WD_ALIGN_PARAGRAPH.LEFT)

    # ---------------- 图片 ----------------
    def _text_width_cm(self):
        c = self.config
        return (21.0 - float(c.get("margin_left_cm", 2.8))
                - float(c.get("margin_right_cm", 2.6)))

    def _fit_images(self, para):
        """把宽出**版心**的图片等比缩回来（同时改锚定框 wp:extent 和图片自身 a:ext）。"""
        max_cx = int(round(self._text_width_cm() * _EMU_PER_CM))
        n = 0
        for drawing in para._p.iter(qn("w:drawing")):
            ext = drawing.find(".//" + _WP_NS + "extent")
            if ext is None:
                continue
            try:
                cx = int(ext.get("cx"))
            except (TypeError, ValueError):
                continue
            if cx <= max_cx or cx <= 0:
                continue
            ratio = float(max_cx) / cx
            for el in list(drawing.iter(_WP_NS + "extent")):
                self._scale_extent(el, ratio)
            for xfrm in drawing.iter(_DRAWING_NS + "xfrm"):
                inner = xfrm.find(_DRAWING_NS + "ext")
                if inner is not None:
                    self._scale_extent(inner, ratio)
            n += 1
        return n

    @staticmethod
    def _scale_extent(el, ratio):
        try:
            el.set("cx", str(int(int(el.get("cx")) * ratio)))
            el.set("cy", str(int(int(el.get("cy")) * ratio)))
        except (TypeError, ValueError):
            pass

    def _format_image_para(self, para):
        """图片/图形段：居中、去缩进、行距放开到单倍，并把超宽图片缩回版心。

        正文用的是固定行距（28 磅），大图会被行高裁掉，所以图片段必须放开行距。
        """
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._clear_indent(para, also_left=True)
        pf = para.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0                     # 单倍行距，避免裁图
        return self._fit_images(para)

    def _format_figcaption(self, para):
        """图题：居中、小四、无缩进。"""
        self._apply_font_to_para(para, self.config.get("body_font", BODY_FONT),
                                 self.config.get("figcaption_size", 12))
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._clear_indent(para, also_left=True)
        self._set_line_spacing(para, self.config.get("line_spacing", 28))

    # ---------------- 页码 / 页面 / docGrid ----------------
    def _add_page_number_field(self, para, align):
        para.alignment = align
        pf = self.config.get("page_number_font", PAGE_NUMBER_FONT)
        ps = self.config.get("page_number_size", 14)
        r0 = para.add_run("— ")
        self._set_run_font(r0, pf, ps)
        r = para.add_run()
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        r._r.append(f1); r._r.append(it); r._r.append(f2)
        self._set_run_font(r, pf, ps)
        r2 = para.add_run(" —")
        self._set_run_font(r2, pf, ps)

    def _apply_page_setup(self, doc):
        c = self.config
        for sec in doc.sections:
            sec.page_width = Cm(21.0)
            sec.page_height = Cm(29.7)
            sec.top_margin = Cm(c.get("margin_top_cm", 3.7))
            sec.bottom_margin = Cm(c.get("margin_bottom_cm", 3.5))
            sec.left_margin = Cm(c.get("margin_left_cm", 2.8))
            sec.right_margin = Cm(c.get("margin_right_cm", 2.6))
            sec.footer_distance = Cm(c.get("footer_distance_cm", 2.8))
            grid = OxmlElement("w:docGrid")
            grid.set(qn("w:type"), "linesAndChars")
            grid.set(qn("w:linePitch"), str(int(c.get("grid_line_pitch", 580))))
            grid.set(qn("w:charSpace"), str(int(c.get("grid_char_space", 0))))
            self._place(sec._sectPr, grid, _AFTER_DOCGRID)
        doc.settings.odd_and_even_pages_header_footer = True
        for sec in doc.sections:
            for foot, align in ((sec.footer, WD_ALIGN_PARAGRAPH.RIGHT),
                                (sec.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT)):
                fp = foot.paragraphs[0] if foot.paragraphs else foot.add_paragraph()
                fp.text = ""
                self._add_page_number_field(fp, align)

    # ---------------- 主流程 ----------------
    def format_document(self, in_path, out_path):
        c = self.config
        tmp = self.convert_to_docx(in_path) if os.path.isfile(in_path) else in_path
        doc = Document(tmp)
        # 1) 文本规范化（安全，保留图片/超链接段）
        if c.get("normalize_punctuation", True) or c.get("remove_spaces", True):
            cnt = skip = 0
            for p in doc.paragraphs:
                nt = N.normalize_text(p.text, c)
                if nt != p.text:
                    if self._set_para_text_safe(p, nt):
                        cnt += 1
                    else:
                        skip += 1
            for tb in doc.tables:              # 表格单元格内的文字同样做规范化
                for row in tb.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            nt = N.normalize_text(p.text, c)
                            if nt != p.text:
                                if self._set_para_text_safe(p, nt):
                                    cnt += 1
                                else:
                                    skip += 1
            self.log("  > 已规范化 %d 段（去空格/统一标点）%s。" % (
                cnt, ("；%d 段含图片/链接未覆盖" % skip) if skip else ""))
        # 1.5) 全文整段成文日期去虚位（成文日期不一定在文末）
        for p in doc.paragraphs:
            dt = (p.text or "").strip()
            if re.match(r"^\d{4}年\d{1,2}月\d{1,2}日$", dt):
                nd = _cn_date(dt)
                if nd != dt:
                    self._set_para_text_safe(p, nd)
        # 1.7) ★问题1彻底★ 全文清除字符间距/字符缩放（含样式与段落标记）
        self._strip_char_spacing(doc)
        # 2) ★A1★ 展开自动编号（并移除 numPr）
        S.expand_auto_numbering(doc, self.log)
        # 3) 去空段（保媒体）
        self._remove_blank_paras(doc)
        # 3.5) ★字距过大★ 清除段内软回车（否则两端对齐会把该行撑满、字距被拉大）
        brn = self._strip_soft_breaks(doc)
        if brn:
            self.log("  > 已清除 %d 处段内软回车（防字距被撑大）。" % brn)
        # 4) 标题识别 + 逐段套用（版头稍后插到最前，避免被误判/重复排版）
        blocks = list(self._iter_blocks(doc))
        title_input = [b for b in blocks if id(b._p) not in self._mast_ids]
        title_paras, subtitle_paras = S.find_title(title_input, c)
        self.last_title = re.sub(r"[。！？；，、：]+$", "",
                                 ((title_paras[0].text or "").strip()) if title_paras else "")
        for p in title_paras:
            self._format_title(p)
        for p in subtitle_paras:
            self._format_subtitle(p)
        tset = {id(p._p) for p in title_paras} | {id(p._p) for p in subtitle_paras}
        date_para, issuer_para = self._find_trailing_signature(blocks)


        recipient_done = False
        seen_heading = False
        prev_image = False
        img_on = bool(c.get("enable_image_formatting", True))
        img_n = cap_n = shrunk = 0
        for bi, blk in enumerate(blocks):
            if not isinstance(blk, Paragraph):
                continue
            if id(blk._p) in tset or id(blk._p) in self._mast_ids:
                continue
            t = (blk.text or "").strip()
            # ★图片★ 含图片/图形的段落：居中、去缩进、行距放开（否则固定行距会裁图）、超宽缩回版心
            if img_on and blk._p.xpath(_IMAGE_XPATH):
                shrunk += self._format_image_para(blk)
                img_n += 1
                prev_image = True
                continue
            # ★图题★ 紧跟在图片下面、以“图1/图一”开头的短段：居中、小四
            if (img_on and prev_image and t and len(t) <= 40
                    and _RE_FIGCAP.match(t)):
                self._format_figcaption(blk)
                cap_n += 1
                prev_image = False
                continue
            prev_image = False
            if not t:
                continue
            if blk is date_para or blk is issuer_para:
                continue
            lvl = S.heading_level(t)
            if lvl:
                seen_heading = True
            if c.get("enable_attachment_formatting", True) and self._is_attachment(blk):
                self._format_attachment_note(blk)
                continue
            if re.match(r"^附件\d*$", t):   # 附件另页标记：3号黑体顶格（不加粗，黑体已够重）
                self._apply_font_to_para(blk, self.config.get("h1_font", H1_FONT),
                                         self.config.get("body_size", 16))
                self._clear_indent(blk, also_left=True)
                self._set_line_spacing(blk, self.config.get("line_spacing", 28))
                blk.alignment = WD_ALIGN_PARAGRAPH.LEFT
                continue
            is_recipient = (not lvl and not recipient_done and not seen_heading
                            and t.endswith("：") and len(t) <= 50)
            if is_recipient:
                self._format_recipient(blk)
                recipient_done = True
                continue
            if lvl == 2 and self._try_h2_inline_body(blk):
                continue
            if lvl:
                self._format_heading(blk, lvl)
            else:
                self._format_body(blk)
        if img_n or cap_n:
            msg = "  > 已处理 %d 个图片段（居中/去缩进/放开行距" % img_n
            msg += "，其中 %d 个超宽已缩回版心" % shrunk if shrunk else ""
            msg += "）"
            if cap_n:
                msg += "，%d 个图题已居中" % cap_n
            self.log(msg + "。")
        # 5) 版头（三档）：正文套用完成后再插到最前
        style = c.get("header_style", HEADER_STYLE_NONE)
        if style != HEADER_STYLE_NONE:
            self._build_header(doc, style)
        # 6) 落款
        if issuer_para or date_para:
            self._format_signature(date_para, issuer_para)
        else:
            self._append_signature_block(doc)
        # 7) 版记
        self._build_footer_notes(doc)
        # 8) 表格
        self._format_tables(doc)
        # 9) 页面/docGrid/页码
        self._apply_page_setup(doc)
        # 9.5) ★字间距过大★ 关闭段落“对齐到网格”（防字符网格撑大字距）
        self._disable_grid_snap(doc)
        # 10) 自检 + 保存
        report = self.verify_document(doc, bool(title_paras))
        doc.save(out_path)
        self.log("  > 已保存：%s" % out_path)
        return report

    def _iter_blocks(self, doc):
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)

    def _remove_blank_paras(self, doc):
        if not self.config.get("remove_blank_lines", True):
            return
        removed = 0
        for p in list(doc.paragraphs):
            if p.text.strip():
                continue
            if p._p.xpath(_BLANK_MEDIA_XPATH):
                continue
            pPr = p._p.pPr
            if pPr is not None and (pPr.find(qn("w:sectPr")) is not None
                                    or pPr.find(qn("w:framePr")) is not None):
                continue
            p._p.getparent().remove(p._p)
            removed += 1
        if removed:
            self.log("  > 去除 %d 个空段。" % removed)

    def _build_header(self, doc, style):
        c = self.config
        org = c.get("org_name", "")
        big = (style == HEADER_STYLE_BIG)
        if not doc.paragraphs:
            doc.add_paragraph("")
        anchor = doc.paragraphs[0]
        # 大红头顶部留白（发文机关标志上边缘距版心上边缘≈35mm，用整行空段近似）
        if big:
            for _ in range(max(0, int(c.get("mast_top_blank_lines", 2)))):
                anchor._p.addprevious(OxmlElement("w:p"))
        new_paras = []
        if org:
            mp = anchor.insert_paragraph_before("")
            run = mp.add_run(org + ("文件" if big else ""))
            self._set_run_font(run, c.get("mast_font", TITLE_FONT),
                               c.get("mast_size_big", 26) if big else c.get("mast_size_small", 22),
                               color=RGBColor(*MAST_RED_RGB))
            mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_paras.append(mp)
        docno = c.get("doc_number", "")
        signer = c.get("signer", "")
        if signer:
            # 上行文：文号左空一字 + 右制表位“签发人：”，标签仿宋、姓名楷体，同行
            line = anchor.insert_paragraph_before("")
            try:
                line.paragraph_format.tab_stops.add_tab_stop(Cm(15.6), WD_TAB_ALIGNMENT.RIGHT)
            except Exception:
                pass
            if docno:
                rd = line.add_run(docno)
                self._set_run_font(rd, c.get("body_font", BODY_FONT), c.get("body_size", 16))
                line.add_run("\t")
            else:
                line.add_run("\t")
            rl = line.add_run("签发人：")
            self._set_run_font(rl, c.get("body_font", BODY_FONT), c.get("body_size", 16))
            rn = line.add_run(signer)
            self._set_run_font(rn, c.get("h2_font", H2_FONT), c.get("body_size", 16))
            self._set_left_indent(line, 1)
            self._set_right_indent(line, 1)
            new_paras.append(line)
        elif docno:
            dp = anchor.insert_paragraph_before("")
            run = dp.add_run(docno)
            self._set_run_font(run, c.get("body_font", BODY_FONT), c.get("body_size", 16))
            dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_paras.append(dp)
        if new_paras:
            self._para_border(new_paras[-1], bottom_sz=12, color="C00000")
            self._mast_ids = {id(p._p) for p in new_paras}
        self.log("  > 已插入版头样式：%s" % style)

    # ---------------- 自检（C2）----------------
    def verify_document(self, doc, title_found):
        checks = []

        def add(name, ok, detail=""):
            checks.append((name, bool(ok), detail))

        c = self.config
        add("主标题识别", title_found, "已识别到大标题" if title_found else "未识别到（可能正文或首段设置）")
        sec = doc.sections[0]
        add("页面 A4", round(sec.page_width.cm, 1) == 21.0 and round(sec.page_height.cm, 1) == 29.7)
        add("页边距(天37/订28)", abs(sec.top_margin.cm - c.get("margin_top_cm", 3.7)) < 0.05
            and abs(sec.left_margin.cm - c.get("margin_left_cm", 2.8)) < 0.05)
        grid = sec._sectPr.find(qn("w:docGrid"))
        add("22行/28字网格(docGrid)", grid is not None and grid.get(qn("w:type")) == "linesAndChars")
        body_ok = any(
            r._r.find(qn("w:rPr")) is not None
            and r._r.find(qn("w:rPr")).find(qn("w:rFonts")) is not None
            and r._r.find(qn("w:rPr")).find(qn("w:rFonts")).get(qn("w:eastAsia")) == c.get("body_font", BODY_FONT)
            for p in doc.paragraphs for r in p.runs)
        add("正文字体 %s(eastAsia)" % c.get("body_font", BODY_FONT), body_ok)
        has_outline = any(
            p._p.pPr is not None and p._p.pPr.find(qn("w:outlineLvl")) is not None
            for p in doc.paragraphs)
        add("标题大纲级别(outlineLvl)", has_outline)
        leftover = sum(1 for p in doc.paragraphs
                       if p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None)
        add("自动编号已处理(A1)", leftover == 0,
            "仍带自动编号 %d 段(项目符号/未判定)" % leftover if leftover else "")
        if c.get("enable_footer_notes", True) and (c.get("cc") or c.get("print_org") or c.get("print_date")):
            all_text = "\n".join((p.text or "") for p in doc.paragraphs)
            add("版记(抄送/印发)(B3)", ("抄送：" in all_text) and ("印发" in all_text))
        return checks


_RESERVED = {"CON", "PRN", "AUX", "NUL"} | {("COM%d" % i) for i in range(1, 10)} | {("LPT%d" % i) for i in range(1, 10)}


def build_filename(title, date_text=None, fallback="公文排版"):
    """★C1★ 导出文件名 = 公文标题 + 系统当天日期 YYYYMMDD。"""
    import datetime
    base = (title or "").strip() or fallback
    base = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "", base)
    base = re.sub(r"[。！？；，、：]+$", "", base)   # 去掉标题尾部句读
    base = re.sub(r"\s+", "", base)
    base = base[:40]
    base = re.sub(r"[.]+$", "", base).lstrip(".")
    if not base:
        base = fallback
    if base.upper() in _RESERVED or re.match(r"^(CON|PRN|AUX|NUL|COM|LPT)\d?$", base.upper()):
        base = "_" + base
    if not date_text:
        date_text = datetime.date.today().strftime("%Y%m%d")
    date_text = re.sub(r"[^0-9]", "", date_text) or datetime.date.today().strftime("%Y%m%d")
    return "%s%s.docx" % (base, date_text)
