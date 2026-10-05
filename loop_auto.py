# -*- coding: utf-8 -*-
"""自动测试循环：每轮随机生成脏公文→引擎排版→GB/T核查，直到连续5轮0违规。"""
import os
import sys
import random
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import gwps_engine as E
import gbt_check
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

CASES = os.path.join(HERE, "_cases")


def set_charspacing(run, val):
    """给 run 设置放大的字间距（模拟问题1的脏输入）。"""
    rPr = run._r.get_or_add_rPr()
    sp = rPr.makeelement(qn("w:spacing"), {qn("w:val"): str(val)})
    rPr.append(sp)

WZ = ["通知", "意见", "报告", "请示", "批复", "决定", "通报", "函", "方案", "规划"]
TOPICS = ["安全生产", "应急管理", "消防安全", "道路交通安全", "防灾减灾", "特种设备安全", "危险化学品安全", "建筑施工安全"]
UNITS = ["各市人民政府", "省直各部门", "各直属机构", "各有关中央驻X单位", "各大企业", "各高等院校", "各市安委会"]
ORG_TAIL = ["人民政府安全生产委员会", "人民政府办公室", "应急管理厅", "安全生产委员会办公室"]
ACTS = ["压实责任", "细化措施", "狠抓落实", "强化督导", "排查隐患", "健全机制", "完善制度", "加强培训", "严格执法", "闭环管理"]
NOUNS = ["人民群众生命财产安全", "安全发展理念", "双重预防机制", "风险分级管控", "隐患排查治理", "应急救援能力", "安全生产责任制"]

SENT = [
    "要坚持以人民为中心的发展思想，树牢安全发展理念，把保护人民生命安全摆在首位，弘扬生命至上、安全第一。",
    "各地区各有关部门要切实提高站位，深刻认识这项工作的极端重要性，统筹发展和安全，坚决防范遏制各类事故发生。",
    "要健全责任体系，压实属地、部门和企业三方责任，做到党政同责、一岗双责、齐抓共管、失职追责。",
    "要深化重点行业领域安全整治，聚焦风险高的环节，开展全覆盖排查，建立台账，限期整改，闭环销号。",
    "要强化科技和人才支撑，推进信息化建设，运用大数据提升风险监测预警能力，加强专业人才培养。",
    "要注重长效机制建设，把好的经验做法固化为制度，动态更新风险清单，防止问题反弹、隐患回潮。",
    "要加强宣传引导，推动安全宣传进企业、进农村、进社区、进学校、进家庭，提升全民安全意识和自救互救能力。",
    "要严格督查考核，将相关工作纳入综合评价和绩效考核，对工作不力的严肃约谈问责，确保各项任务落地见效。",
]


def body(rng, n=1):
    out = []
    for _ in range(n):
        out.append(rng.choice(SENT))
    return "".join(out)


def inject(rng, s):
    """往正文随机注入待修复的问题（每条都能被引擎规范化，确保可收敛）。"""
    r = rng.random()
    if r < 0.10:
        i = s.find("，")
        if i > 0:
            return s[:i] + "  " + s[i:]              # 双半角空格
    if r < 0.18:
        return s.replace("。", "。　", 1)             # 全角空格
    if r < 0.26:
        return s.replace("要", "要\u200b", 1)         # 零宽符
    if r < 0.34:
        return s.replace("要", "要\u00a0", 1)         # 不间断空格
    if r < 0.42:
        return s.replace("，", "\t", 1)               # 制表符
    if r < 0.52:
        kw = rng.choice(NOUNS)
        return s.replace(kw, '"' + kw + '"', 1)       # 直引号（成对）
    if r < 0.60:
        return s.replace("，", ",", 1).replace("。", ".", 1)   # 半角标点
    if r < 0.68:
        return s + "等等...要抓紧。"                    # 省略号
    if r < 0.76:
        return s + "资料汇编2020--2025。"               # 双连字符
    if r < 0.84:
        return s + "电话0571-8888-1234。"               # 半角连字符(须保留)
    if r < 0.92:
        return s + "会议12:30召开，预算1,000万元。"       # 时间/千分位(须保留)
    return s


def make_png(path, w=1600, h=500):
    """造一张测试用 PNG。"""
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle([6, 6, w - 7, h - 7], outline=(122, 15, 34), width=8)
    for i in range(1, 5):
        x = int(w * i / 5.0)
        d.line([x, 30, x, h - 30], fill=(212, 175, 55), width=6)
    im.save(path)
    return path


def add_p(doc, text, center=False, soft=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    if soft:
        r.add_break()
        p.add_run(soft)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def gen_doc(rng):
    doc = Document()
    add_p(doc, "　　" + "关于做好本年度" + rng.choice(TOPICS) + "工作的" + rng.choice(WZ) + "。")
    add_p(doc, " " + "，".join(rng.sample(UNITS, rng.randint(2, 4))) + ":")
    p = add_p(doc, inject(rng, body(rng, 2)))
    set_charspacing(p.runs[0], rng.choice([40, 60, 80]))   # ★问题1★ 正文放大字间距
    h1n = rng.randint(4, 7)
    for i in range(h1n):
        add_p(doc, cn(i + 1) + "、" + rng.choice(["总体要求", "主要任务", "工作措施", "责任分工", "保障要求", "进度安排", "监督检查", "长效机制"]))
        for j in range(rng.randint(1, 3)):
            sub = "（" + cncn(j + 1) + "）"
            if rng.random() < 0.5:  # 引导句式
                add_p(doc, sub + rng.choice(["提高认识", "压实责任", "强化监管", "夯实基础"]) + "。" + inject(rng, body(rng, 1)))
            else:  # 独立二级 + 下级
                add_p(doc, sub + rng.choice(["加强统筹", "严格执法", "科技支撑", "宣传培训"]))
                add_p(doc, str(rng.randint(1, 3)) + "." + rng.choice(["细化方案", "明确分工", "建立台账", "闭环整改"]))
                add_p(doc, "（" + str(rng.randint(1, 3)) + "）" + inject(rng, body(rng, 1)))
        add_p(doc, inject(rng, body(rng, 2)))
    # 表格
    rows = rng.randint(3, 6)
    add_p(doc, "附表　工作安排表")
    t = doc.add_table(rows=rows, cols=3)
    data = [("阶段", "时间", "任务"),
            ("部署", "2026年7月", "制定方案,发动部署"),
            ("排查", "2026年8-9月", "全面排查,建立台账"),
            ("整治", "2026年10月", "集中整改,挂牌督办"),
            ("巩固", "2026年11月", "建章立制,总结评估"),
            ("回头看", "2026年12月", "防止反弹,闭环管理")]
    for i in range(rows):
        for j in range(3):
            t.cell(i, j).text = data[i][j]
    # 图片：一张故意超版心宽(20cm > 15.6cm，应被缩回)，一张正常宽；各带图题
    png = make_png(os.path.join(HERE, "_test_img.png"))
    add_p(doc, "附：工作安排示意图")
    doc.add_picture(png, width=Cm(20))
    add_p(doc, "图1 工作安排示意图")
    doc.add_picture(png, width=Cm(9))
    add_p(doc, "图2 责任分工示意图")
    # 软回车/制表 段
    add_p(doc, "第一，常态化排查", soft="对重点区域开展不间断巡查，发现问题立行立改。")
    # 附件说明
    items = rng.sample(["重点任务分工表", "隐患排查清单", "督导检查办法", "考核问责细则"], rng.randint(1, 3))
    if rng.random() < 0.6:  # ★问题2★ 附件序号间无半角空格
        note = "附件：" + "".join("%d.%s" % (k + 1, it) for k, it in enumerate(items))
    else:
        note = "附件：" + " ".join("%d.%s" % (k + 1, it) for k, it in enumerate(items))
    add_p(doc, note)
    # 署名 + 成文日期（随机虚位）
    add_p(doc, rng.choice(ORG_TAIL))
    mo = rng.choice(["0%d" % rng.randint(1, 9), "%d" % rng.randint(10, 12)])
    dd = rng.choice(["0%d" % rng.randint(1, 9), "%d" % rng.randint(10, 28)])
    add_p(doc, "2026年%s月%s日" % (mo, dd))
    add_p(doc, "（此件公开发布。联系人：%s；联系电话：0571-8888-1234）" % rng.choice(["张 伟", "王 芳", "李 强"]))
    # 附件另页
    doc.add_page_break()
    add_p(doc, "附件1")
    pt = add_p(doc, rng.choice(items))
    set_charspacing(pt.runs[0], rng.choice([40, 60]))   # ★问题2★ 附件标题字间距
    add_p(doc, body(rng, 2))
    doc.add_page_break()
    add_p(doc, "附件2")
    pt2 = add_p(doc, rng.choice(items))
    set_charspacing(pt2.runs[0], rng.choice([40, 60]))
    add_p(doc, body(rng, 2))
    # 版记
    add_p(doc, "抄送：" + "，".join(rng.sample(UNITS, rng.randint(2, 4))) + "。")
    add_p(doc, rng.choice(ORG_TAIL) + "    2026年8月6日印发")
    # 兜底加长
    while sum(len(p.text) for p in doc.paragraphs) < 5200:
        add_p(doc, inject(rng, body(rng, 2)))
    return doc


def cn(n):
    return "一二三四五六七八九十"[n - 1] if 1 <= n <= 10 else str(n)


def cncn(n):
    return "一二三四五六七八九"[n - 1] if 1 <= n <= 9 else str(n)


def run_rounds(max_rounds=20, need_clean=5):
    shutil.rmtree(CASES, ignore_errors=True)
    os.makedirs(CASES, exist_ok=True)
    clean_streak = 0
    round_id = 0
    last_viol = []
    while round_id < max_rounds and clean_streak < need_clean:
        round_id += 1
        rng = random.Random(1000 + round_id)
        doc = gen_doc(rng)
        chars = sum(len(p.text) for p in doc.paragraphs)
        inp = os.path.join(CASES, "r%d.docx" % round_id)
        outp = os.path.join(CASES, "r%d_out.docx" % round_id)
        doc.save(inp)
        try:
            E.WordProcessor(config={}, log=lambda m: None).format_document(inp, outp)
            viol = gbt_check.check(outp)
        except Exception as e:
            viol = ["EXCEPTION: " + repr(e)]
        if chars < 5000:
            viol = viol + ["生成不足5000字: %d" % chars]
        tag = "CLEAN" if not viol else ("VIOL %d" % len(viol))
        print("ROUND %d chars=%d -> %s" % (round_id, chars, tag))
        for x in viol[:8]:
            print("     -", x)
        if not viol:
            clean_streak += 1
            last_viol = []
        else:
            clean_streak = 0
            last_viol = viol
    print("\n#### 结果：完成轮=%d 连续洁净=%d ####" % (round_id, clean_streak))
    if clean_streak >= need_clean:
        print("CONVERGED")
    else:
        print("NOT_CONVERGED，最后违规：")
        for x in last_viol:
            print("   -", x)


if __name__ == "__main__":
    run_rounds()
